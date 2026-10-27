import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.abspath(".")
SCREENSHOTS_DIR = os.path.join(BASE_DIR, "screenshots")
OUTPUT_DOCX = os.path.join(BASE_DIR, "StudentHub_Responsive_Design_Report.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="EFF6FF", border_hex="2563EB"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.font.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(29, 78, 216)
    
    run_txt = p.add_run(text)
    run_txt.font.size = Pt(9.5)
    run_txt.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text, max_lines=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    lines = code_text.splitlines()
    if max_lines and len(lines) > max_lines:
        displayed_lines = lines[:max_lines]
        displayed_lines.append(f"... [{len(lines) - max_lines} more lines in actual file]")
    else:
        displayed_lines = lines
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    
    for idx, line in enumerate(displayed_lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line if line.strip() else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def read_file_safe(rel_path):
    full_path = os.path.join(BASE_DIR, rel_path.replace("/", "\\"))
    if os.path.exists(full_path):
        with open(full_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    return ""

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
        # Header
        hdr = s.header
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("StudentHub — Responsive Web Portal Engineering Report")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        ftr = s.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("CSS Grid & Flexbox Mobile-First Implementation")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def create_document():
    doc = docx.Document()
    add_header_footer(doc)
    
    # Title Section
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(2)
    trun = title_p.add_run("StudentHub Academic Portal")
    trun.font.name = "Segoe UI"
    trun.font.size = Pt(26)
    trun.font.bold = True
    trun.font.color.rgb = RGBColor(37, 99, 235)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    srun = sub_p.add_run("Responsive UI/UX Engineering Report: CSS Grid, Flexbox, Mobile-First Architecture, Source Code & Visual Outputs")
    srun.font.name = "Segoe UI"
    srun.font.size = Pt(12)
    srun.font.color.rgb = RGBColor(100, 116, 139)
    
    # Metadata Table
    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    col_w = [Inches(2.2), Inches(4.3)]
    
    meta_data = [
        ("Project Scope", "Design responsive layouts for Home, About, Registration, Dashboard, and Events (+ Full Portal)"),
        ("Layout Engine", "CSS Grid (2, 3, 4-column auto-fit grids) and Flexbox (navigation, hero, card bars, headers)"),
        ("Responsive Strategy", "Mobile-First (Base: <768px, Tablet: 768px-1024px, Desktop: >=1024px)"),
        ("Documentation Scope", "Full source code, architectural breakdown, and dual Desktop/Mobile output screenshots")
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        row = meta_tbl.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = col_w[0]
        c1.width = col_w[1]
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGBColor(30, 41, 59)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(51, 65, 85)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section 1: Executive Overview
    h1 = doc.add_heading("1. Executive Summary & Design System Architecture", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph()
    p.add_run(
        "StudentHub is a comprehensive college web portal designed to provide students with intuitive, real-time access "
        "to academic curricula, assignments, grades, attendance tracking, and campus events. The primary design objective was to "
        "re-architect the entire portal using modern CSS Grid and Flexbox with a strict mobile-first paradigm, completely "
        "eliminating viewport overflow, raw unstyled layout stretches, and accessibility barriers."
    )
    
    add_callout(
        doc,
        "Core Technical Pillars",
        "• Mobile-First Responsive CSS: Base styles target smartphone viewports (<768px), progressively enhanced via media queries for tablets (min-width: 768px) and desktop displays (min-width: 1024px).\n"
        "• CSS Grid System: Powers multi-column card dashboards, feature showcases, statistical metric cards (auto-fit minmax), and multi-field registration forms.\n"
        "• Flexbox Engine: Powers sticky header navigation, off-canvas drawer toggles, hero action alignments, badge status chips, and responsive data tables.\n"
        "• WCAG 2.1 Accessibility: Skip-navigation anchors, semantic landmark tags (<header>, <nav>, <main>, <aside>, <footer>), and 44px minimum touch targets."
    )
    
    # Section 2: Global Stylesheet & JavaScript
    h1 = doc.add_heading("2. Core Global Architecture: CSS & Navigation Script", level=1)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)
    
    doc.add_heading("2.1 Master Stylesheet (css/style.css)", level=2)
    doc.add_paragraph(
        "The master stylesheet defines design tokens (variables), base typography resets, container utility rules, "
        "CSS Grid card layouts, responsive tables, badge systems, and media queries across all breakpoints."
    )
    css_content = read_file_safe("css/style.css")
    add_code_block(doc, css_content, max_lines=180)
    
    doc.add_heading("2.2 Mobile Navigation Toggle Script (js/main.js)", level=2)
    doc.add_paragraph(
        "A lightweight, vanilla JavaScript module providing smooth hamburger menu toggling, outside-click dismissal, "
        "and keyboard accessibility (Escape key close)."
    )
    js_content = read_file_safe("js/main.js")
    add_code_block(doc, js_content)
    
    # Core Requested Pages List
    PAGES_INFO = [
        {
            "name": "Home Page",
            "file": "index.html",
            "key": "index",
            "grid_flex": "• Flexbox: Split Hero Banner (hero text on left, live portal card on right), sticky brand navigation.\n• CSS Grid: 3-column auto-fit Features Grid, 4-column Statistics Metric Counters, and 3-column Quick Navigation Links.",
            "desc": "The primary landing page introducing StudentHub features, live semester metrics, enrollment CTAs, and university navigation shortcuts.",
            "before_img": None
        },
        {
            "name": "About Page",
            "file": "pages/about.html",
            "key": "about",
            "grid_flex": "• CSS Grid: 2-column Mission & Vision Cards with left accent borders, 6-card Offerings Grid, and 3-column Team Showcase.\n• Flexbox: Team avatar profile chips, badge tags, and responsive footer links.",
            "desc": "Presents university vision, platform mission, student offerings, and the leadership development faculty team.",
            "before_img": None
        },
        {
            "name": "Registration Page",
            "file": "pages/register.html",
            "key": "register",
            "grid_flex": "• CSS Grid: .form-grid-2 transforms fields into 2 equal columns on tablet/desktop while defaulting to 1 column on mobile devices.\n• Flexbox: Centered card container (.auth-wrapper), custom checkbox alignment, and action buttons.",
            "desc": "A student registration portal featuring structured personal/academic inputs, password confirmation, terms acceptance, and accessible validation states.",
            "before_img": None
        },
        {
            "name": "Student Dashboard",
            "file": "pages/dashboard.html",
            "key": "dashboard",
            "grid_flex": "• CSS Grid: 4-column Academic Summary Cards (Enrolled, Tasks, Attendance, CGPA) + 2-Column Desktop Grid (2.2fr main deadlines table + 1fr alerts sidebar).\n• Flexbox: Personalized student header with avatar badge ('AP') and quick shortcut pills.",
            "desc": "The core operational dashboard for enrolled students, organizing pending assignments, attendance benchmarks, alerts, and navigation links.",
            "before_img": None
        },
        {
            "name": "Campus Events Page",
            "file": "pages/events.html",
            "key": "events",
            "grid_flex": "• Flexbox: Filter category bar with interactive pills (Technical, Cultural, Sports, Workshops).\n• CSS Grid: 3-column Event Cards with dual-tier calendar date badges (Day/Month), venue icons, and responsive tabular calendar overview.",
            "desc": "Highlights university workshops, hackathons, guest lectures, and cultural nights with direct registration CTAs.",
            "before_img": None
        },
        {
            "name": "Course Catalog",
            "file": "pages/courses.html",
            "key": "courses",
            "grid_flex": "• CSS Grid: 4 snapshot metrics for semester terms and credits.\n• Flexbox: Table cell alignments, course code badges (CS301-CS306), and direct assignment links.",
            "desc": "Semester 5 curriculum catalog detailing course codes, subject names, credit ratings, and faculty coordinators.",
            "before_img": "before_courses.png"
        },
        {
            "name": "Assignments Tracker",
            "file": "pages/assignment.html",
            "key": "assignment",
            "grid_flex": "• CSS Grid: 4 status cards (Total, Completed, Pending, Overdue).\n• Flexbox: Contextual status badges (Graded, Submitted, Pending, Overdue) and action buttons (Upload, Feedback).",
            "desc": "Academic task submissions dashboard with deadline monitoring, evaluation marks, and homework upload triggers.",
            "before_img": "before_assignment.png"
        },
        {
            "name": "Attendance Record",
            "file": "pages/attendance.html",
            "key": "attendance",
            "grid_flex": "• CSS Grid: 4 top stat cards replacing the raw vertical table (Total Classes, Attended, Absent, Overall 92%).\n• Flexbox: Eligibility badges (>75% mandatory threshold) and subject-wise breakdown.",
            "desc": "Real-time lecture and laboratory attendance monitor with automatic eligibility alerts.",
            "before_img": "before_attendance.png"
        },
        {
            "name": "Academic Results & Grades",
            "file": "pages/result.html",
            "key": "result",
            "grid_flex": "• CSS Grid: 4 summary cards for SGPA (8.7), CGPA (8.5), Credits (42), and Backlog Status.\n• Flexbox: Letter grade pills (A+, A, B+), 10-point grade point calculations, and print button.",
            "desc": "Official semester examination grade transcript and cumulative performance report.",
            "before_img": "before_result.png"
        },
        {
            "name": "Student Profile",
            "file": "pages/profile.html",
            "key": "profile",
            "grid_flex": "• CSS Grid: 2-column Personal Information Grid and 2-column Academic Information Grid.\n• Flexbox: Profile banner with circular avatar ('RS'), student credentials, and action buttons.",
            "desc": "Comprehensive student identity record containing personal address, phone, roll number, and department standings.",
            "before_img": "before_profile.png"
        },
        {
            "name": "Login Authentication Page",
            "file": "pages/login.html",
            "key": "login",
            "grid_flex": "• Flexbox: Centered authentication card, checkbox session remember row, and submission button.\n• CSS: Focus rings and input field states.",
            "desc": "Secure student login interface with password recovery links and seamless redirect to dashboard.",
            "before_img": None
        },
        {
            "name": "Contact & Support Desk",
            "file": "pages/contact.html",
            "key": "contact",
            "grid_flex": "• CSS Grid: 2-column desktop layout (Send Message form on left, Campus Office address on right).\n• Flexbox: Form header and submit CTA.",
            "desc": "Student help desk portal with department inquiry selector and campus location address.",
            "before_img": None
        },
        {
            "name": "Frequently Asked Questions (FAQ)",
            "file": "pages/faq.html",
            "key": "faq",
            "grid_flex": "• Flexbox: Accordion <details>/<summary> interactive widgets with animated toggle indicators (+/-).\n• CSS Grid: Support card banner.",
            "desc": "Campus knowledgebase answering questions regarding course drops, exams, and attendance rules.",
            "before_img": None
        },
        {
            "name": "Feedback & Experience Form",
            "file": "pages/feedback.html",
            "key": "feedback",
            "grid_flex": "• CSS Grid: 2-column inputs for name and email.\n• Flexbox: Radio rating pills (1 to 5 stars) and centered submission wrapper.",
            "desc": "Interactive rating and suggestion form for student portal improvements.",
            "before_img": None
        },
        {
            "name": "University Admin Panel",
            "file": "pages/admin.html",
            "key": "admin",
            "grid_flex": "• CSS Grid: 4 platform telemetry counters (10,248 students, 542 courses, 28 pending, 99.8% uptime) + 3-col quick admin actions.\n• Flexbox: Approval/Rejection action pills and table controls.",
            "desc": "Campus administrative management portal for reviewing student registrations and course capacities.",
            "before_img": None
        }
    ]
    
    # Section 3: Page-by-Page Detailed Documentation
    h1 = doc.add_heading("3. Page-by-Page Code & Output Screenshots", level=1)
    h1.paragraph_format.space_before = Pt(20)
    h1.paragraph_format.space_after = Pt(8)
    
    for idx, page in enumerate(PAGES_INFO, 1):
        h2 = doc.add_heading(f"3.{idx} {page['name']} ({page['file']})", level=2)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(4)
        
        doc.add_paragraph(f"Description: {page['desc']}")
        add_callout(doc, "CSS Grid & Flexbox Techniques", page['grid_flex'], bg_hex="F8FAFC", border_hex="4F46E5")
        
        # If there is a "Before" screenshot from user, show Before vs After
        if page['before_img']:
            before_path = os.path.join(SCREENSHOTS_DIR, page['before_img'])
            if os.path.exists(before_path):
                doc.add_heading("Before vs After Transformation:", level=3)
                p_before = doc.add_paragraph()
                r_b = p_before.add_run("⚠️ Before Fix: Unstyled raw table stretching infinitely across the viewport:\n")
                r_b.font.italic = True
                r_b.font.color.rgb = RGBColor(220, 38, 38)
                doc.add_picture(before_path, width=Inches(5.8))
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
        
        # Output Screenshots: Desktop & Mobile Side by Side
        d_img = os.path.join(SCREENSHOTS_DIR, f"{page['key']}_desktop.png")
        m_img = os.path.join(SCREENSHOTS_DIR, f"{page['key']}_mobile.png")
        
        doc.add_heading("Rendered Output Screenshots (Desktop & Mobile):", level=3)
        
        if os.path.exists(d_img) and os.path.exists(m_img):
            tbl_img = doc.add_table(rows=2, cols=2)
            tbl_img.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_img.autofit = False
            
            c_d_lbl = tbl_img.cell(0, 0)
            c_m_lbl = tbl_img.cell(0, 1)
            c_d_lbl.width = Inches(4.3)
            c_m_lbl.width = Inches(2.2)
            
            p_dl = c_d_lbl.paragraphs[0]
            p_dl.paragraph_format.space_after = Pt(2)
            rdl = p_dl.add_run("💻 Desktop View (1280px)")
            rdl.font.bold = True
            rdl.font.size = Pt(8.5)
            rdl.font.color.rgb = RGBColor(37, 99, 235)
            
            p_ml = c_m_lbl.paragraphs[0]
            p_ml.paragraph_format.space_after = Pt(2)
            rml = p_ml.add_run("📱 Mobile View (412px)")
            rml.font.bold = True
            rml.font.size = Pt(8.5)
            rml.font.color.rgb = RGBColor(16, 185, 129)
            
            c_d_img = tbl_img.cell(1, 0)
            c_m_img = tbl_img.cell(1, 1)
            c_d_img.width = Inches(4.3)
            c_m_img.width = Inches(2.2)
            
            c_d_img.paragraphs[0].paragraph_format.space_after = Pt(0)
            c_d_img.paragraphs[0].add_run().add_picture(d_img, width=Inches(4.1))
            
            c_m_img.paragraphs[0].paragraph_format.space_after = Pt(0)
            c_m_img.paragraphs[0].add_run().add_picture(m_img, width=Inches(2.0))
            
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
        elif os.path.exists(d_img):
            doc.add_picture(d_img, width=Inches(6.0))
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
            
        # Source Code
        doc.add_heading("Complete Source Code:", level=3)
        code_text = read_file_safe(page['file'])
        add_code_block(doc, code_text)
        
    # Section 4: Testing & Verification Matrix
    h1 = doc.add_heading("4. Responsive Layout Verification Matrix", level=1)
    h1.paragraph_format.space_before = Pt(20)
    h1.paragraph_format.space_after = Pt(8)
    
    matrix_tbl = doc.add_table(rows=6, cols=4)
    matrix_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    matrix_tbl.autofit = False
    
    headers = ["Feature / Page", "CSS Grid Implementation", "Flexbox Implementation", "Mobile Responsiveness"]
    for i, h in enumerate(headers):
        cell = matrix_tbl.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    rows_data = [
        ("Home (index.html)", "3-col Features & 4-col Stats", "Hero Banner & Header Nav", "Full column collapse (1 col)"),
        ("About (about.html)", "2-col Mission & 3-col Team", "Avatar badges & Footer", "Vertical stacking on mobile"),
        ("Registration", ".form-grid-2 2-col inputs", "Centered card & Submit", "100% width mobile touch targets"),
        ("Dashboard", "2-col 70/30 grid + 4 Stat cards", "Avatar bar & Deadlines", "Single column stacked flow"),
        ("Events (events.html)", "3-col Event cards + Cal Table", "Filter Pills & Date badges", "Scrollable table + wrapped chips")
    ]
    
    for r_idx, r_data in enumerate(rows_data, 1):
        for c_idx, val in enumerate(r_data):
            cell = matrix_tbl.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 0:
                r.font.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    
    # Save Document
    doc.save(OUTPUT_DOCX)
    print(f"Document successfully created: {OUTPUT_DOCX} ({os.path.getsize(OUTPUT_DOCX)} bytes)")

if __name__ == "__main__":
    create_document()
